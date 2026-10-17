"""Deadline Sentinel 单元测试。

运行：python -m unittest discover tests -v
或：  python tests/test_sentinel.py

覆盖：读入解析、任务识别、时间解析、任务组装（去重/更正/待确认）、
      导出（CSV/JSON）、AI 模块（用注入的假传输层离线测试，不联网）。
"""

import io
import json
import os
import sys
import unittest
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from sentinel.read_chat import parse_line, parse_lines, read_chat_file
from sentinel.extract import looks_like_task, topic_of, topics_of, confidence_of
from sentinel.ddl import parse_due, parse_due_detail, status_of, remaining_text
from sentinel.tasks import build_tasks, summarize, sort_tasks
from sentinel import export as exporter
from sentinel import ai_helper


ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REF = datetime(2026, 10, 14, 9, 0)   # 2026-10-14 是周三


def sample_path(name="sample_chat.txt"):
    return os.path.join(ROOT, "data", name)


class TestReadChat(unittest.TestCase):
    def test_parse_ok(self):
        msg, reason = parse_line("2026-10-14 09:12 | 助教-李 | 数据结构 | 实验三报告下周日交")
        self.assertEqual(reason, "")
        self.assertEqual(msg["sender"], "助教-李")
        self.assertEqual(msg["chat"], "数据结构")

    def test_parse_fullwidth_separator(self):
        msg, reason = parse_line("2026-10-14 09:12 ｜ 助教-李 ｜ 数据结构 ｜ 实验三报告下周日交")
        self.assertIsNotNone(msg)
        self.assertEqual(msg["chat"], "数据结构")

    def test_parse_bad_reasons(self):
        _, r1 = parse_line("")
        self.assertEqual(r1, "空行")
        _, r2 = parse_line("格式不对没有竖线")
        self.assertIn("字段数", r2)
        _, r3 = parse_line("坏时间 | 人 | 群 | 内容")
        self.assertIn("时间格式", r3)

    def test_parse_lines_counts_problems(self):
        text = "2026-10-14 09:12 | A | B | 作业下周日交\n\n这一行是坏的\n"
        messages, problems = parse_lines(text)
        self.assertEqual(len(messages), 1)
        self.assertEqual(len(problems), 1)
        self.assertEqual(problems[0]["line_no"], 3)

    def test_read_sample(self):
        messages, problems = read_chat_file(sample_path())
        self.assertGreaterEqual(len(messages), 15)
        self.assertEqual(problems, [])


class TestExtract(unittest.TestCase):
    def test_task_yes(self):
        self.assertTrue(looks_like_task("实验三报告下周日晚上12点前交"))

    def test_noise_no(self):
        self.assertFalse(looks_like_task("收到收到"))
        self.assertFalse(looks_like_task("[广告] 考研资料大促销，点击链接"))
        self.assertFalse(looks_like_task("这题谁会做？在线等"))

    def test_topic_variants(self):
        self.assertEqual(topic_of("实验三报告下周日交"), "实验三")
        self.assertEqual(topic_of("第3次作业下周一交"), "第3次作业")
        self.assertEqual(topic_of("课程设计选题月底前报"), "课程设计")

    def test_multi_topics(self):
        ts = topics_of("实验一和实验二的报告都要交")
        self.assertIn("实验一", ts)
        self.assertIn("实验二", ts)

    def test_confidence_vague_lower(self):
        solid = confidence_of("实验三报告下周日晚上12点前交")
        vague = confidence_of("实验三报告大概可能下周交")
        self.assertLess(vague, solid)


class TestDdl(unittest.TestCase):
    def test_absolute(self):
        due, src = parse_due("10月20日 23:59 交", REF)
        self.assertEqual((due.year, due.month, due.day, due.hour, due.minute), (2026, 10, 20, 23, 59))

    def test_absolute_detail_explicit(self):
        d = parse_due_detail("2026-11-06 上午考试", REF)
        self.assertTrue(d["explicit"])
        self.assertFalse(d["ambiguous"])

    def test_next_weekday(self):
        due, src = parse_due("下周日晚上12点前交", REF)
        self.assertEqual((due.month, due.day), (10, 25))
        self.assertEqual((due.hour, due.minute), (23, 59))

    def test_weekday(self):
        due, src = parse_due("本周五课堂小测", REF)
        self.assertEqual((due.month, due.day), (10, 16))

    def test_month_end(self):
        due, src = parse_due("月底前把课程设计选题报一下", REF)
        self.assertEqual((due.month, due.day), (10, 31))

    def test_past_month_rolls_next_year(self):
        due, src = parse_due("3月5日交第二次作业", REF)
        self.assertEqual((due.year, due.month, due.day), (2027, 3, 5))
        detail = parse_due_detail("3月5日交第二次作业", REF)
        self.assertTrue(detail["ambiguous"])  # 没写年份，标为含糊

    def test_unknown_stays_none(self):
        d = parse_due_detail("验收时间待定，另行通知", REF)
        self.assertIsNone(d["due"])
        self.assertEqual(d["due_source"], "无法识别")

    def test_status(self):
        self.assertEqual(status_of(datetime(2026, 10, 14, 8, 0), REF), "已过期")
        self.assertEqual(status_of(datetime(2026, 10, 14, 20, 0), REF), "紧急")
        self.assertEqual(status_of(datetime(2026, 10, 16, 20, 0), REF), "即将")

    def test_remaining_text(self):
        self.assertTrue(remaining_text(datetime(2026, 10, 14, 8, 0), REF).startswith("已过期"))
        self.assertTrue(remaining_text(None, REF) == "未知")


class TestBuildTasks(unittest.TestCase):
    def _tasks(self, now=datetime(2026, 10, 14, 18, 57)):
        messages, _ = read_chat_file(sample_path())
        return build_tasks(messages, now=now)

    def test_correction_supersedes(self):
        tasks = self._tasks()
        dead = [t for t in tasks if t["superseded"]]
        self.assertTrue(any("10月20日" in t["text"] for t in dead))

    def test_unknown_not_fabricated(self):
        tasks = self._tasks()
        # 没有任何日期信息的消息必须保留 due=None，不能编造日期
        unknown = [t for t in tasks if t["due"] is None]
        self.assertTrue(unknown)                      # 至少有一条未知
        for t in unknown:
            self.assertEqual(t["due_source"], "无法识别")
            self.assertIn(t["status"], ("未知", "待确认", "已作废"))
        # 更正类消息（以……为准）应保留未知并标为待确认
        corrections = [t for t in tasks if "为准" in t["text"]]
        self.assertTrue(corrections)
        for t in corrections:
            self.assertIsNone(t["due"])
            self.assertEqual(t["status"], "待确认")

    def test_needs_confirm_flag(self):
        tasks = self._tasks()
        # 含糊 / 未知 / 低置信度的任务应被建议确认
        self.assertTrue(any(t["needs_confirm"] for t in tasks))

    def test_original_message_preserved(self):
        tasks = self._tasks()
        for t in tasks:
            self.assertIn("text", t)
            self.assertTrue(t["text"])          # 原文永远保留
            self.assertIn("msg_time", t)

    def test_stable_ids_unique(self):
        tasks = self._tasks()
        ids = [t["id"] for t in tasks]
        self.assertEqual(len(ids), len(set(ids)))

    def test_summary_keys(self):
        s = summarize(self._tasks())
        for k in ("total", "urgent", "soon", "confirm", "overdue", "superseded"):
            self.assertIn(k, s)


class TestExport(unittest.TestCase):
    def _tasks(self):
        messages, _ = read_chat_file(sample_path())
        return build_tasks(messages, now=datetime(2026, 10, 14, 18, 57))

    def test_csv_roundtrip(self):
        text = exporter.to_csv(self._tasks())
        self.assertIn("截止时间", text)
        self.assertIn("原始消息", text)

    def test_json_has_traceability(self):
        payload = json.loads(exporter.to_json(self._tasks()))
        self.assertTrue(payload)
        first = payload[0]
        for k in ("raw_message", "due_source", "origin", "explicit", "needs_confirm"):
            self.assertIn(k, first)

    def test_no_secret_in_export(self):
        os.environ["API_KEY"] = "sk-SECRET-should-not-leak"
        try:
            blob = exporter.to_json(self._tasks()) + exporter.to_csv(self._tasks())
            self.assertNotIn("sk-SECRET-should-not-leak", blob)
        finally:
            os.environ.pop("API_KEY", None)


class TestAiHelperOffline(unittest.TestCase):
    """用注入的假传输层测试 AI 模块，完全不联网。"""

    ENV = {"BASE_URL": "https://example.invalid/api/v1", "API_KEY": "sk-test", "MODEL": "test-model"}

    def _tasks(self):
        messages, _ = read_chat_file(sample_path())
        return build_tasks(messages, now=datetime(2026, 10, 14, 18, 57))

    def test_disabled_without_env(self):
        suggestions, meta = ai_helper.suggest_for_tasks(self._tasks(), REF, env={})
        self.assertFalse(meta["enabled"])
        self.assertEqual(suggestions, [])
        self.assertIn("未配置", meta["error"])

    def test_selects_only_uncertain(self):
        picked = ai_helper.select_tasks_for_ai(self._tasks())
        self.assertTrue(picked)
        for t in picked:
            self.assertTrue(t["due"] is None or t["ambiguous"] or t["needs_confirm"])

    def test_good_response_applied_and_validated(self):
        tasks = self._tasks()
        picked = ai_helper.select_tasks_for_ai(tasks)
        target = picked[0]

        def fake_transport(payload, cfg, timeout):
            # 同时塞一个不存在的 id，验证会被丢弃
            content = json.dumps([
                {"id": target["id"], "due": "2026-11-01 23:59", "reason": "测试"},
                {"id": "NOT-IN-INPUT", "due": "2026-11-02 23:59", "reason": "伪造"},
            ], ensure_ascii=False)
            return {"choices": [{"message": {"content": content}}]}

        suggestions, meta = ai_helper.suggest_for_tasks(
            tasks, REF, env=self.ENV, transport=fake_transport)
        self.assertTrue(meta["used"])
        self.assertEqual(len(suggestions), 1)          # 伪造 id 被丢弃
        self.assertEqual(suggestions[0]["id"], target["id"])

    def test_bad_json_rejected(self):
        def fake_transport(payload, cfg, timeout):
            return {"choices": [{"message": {"content": "这不是JSON"}}]}
        suggestions, meta = ai_helper.suggest_for_tasks(
            self._tasks(), REF, env=self.ENV, transport=fake_transport)
        self.assertEqual(suggestions, [])
        self.assertTrue(meta["error"])

    def test_out_of_range_date_rejected(self):
        tasks = self._tasks()
        target = ai_helper.select_tasks_for_ai(tasks)[0]

        def fake_transport(payload, cfg, timeout):
            content = json.dumps([{"id": target["id"], "due": "2000-01-01 00:00", "reason": "离谱"}])
            return {"choices": [{"message": {"content": content}}]}

        suggestions, meta = ai_helper.suggest_for_tasks(
            tasks, REF, env=self.ENV, transport=fake_transport)
        self.assertEqual(suggestions, [])   # 日期太早，被校验丢弃

    def test_error_kind_mapping(self):
        def fake_transport(payload, cfg, timeout):
            raise ai_helper.AIError("rate_limit", "被限流")
        suggestions, meta = ai_helper.suggest_for_tasks(
            self._tasks(), REF, env=self.ENV, transport=fake_transport)
        self.assertEqual(suggestions, [])
        self.assertIn("限流", meta["error"])

    def test_mask_key(self):
        self.assertNotIn("sk-abcdefghij", ai_helper.mask_key("sk-abcdefghij"))
        self.assertEqual(ai_helper.mask_key(""), "(未配置)")


class TestWebappFlow(unittest.TestCase):
    """不启服务，直接测分析流程（含 AI 关闭时的路径）。"""

    def test_analyze_without_ai(self):
        import webapp
        text = open(sample_path(), encoding="utf-8").read()
        result = webapp.analyze_text(text, use_ai=False, now=datetime(2026, 10, 14, 18, 57))
        self.assertTrue(result["ok"])
        self.assertIn("summary", result)
        self.assertIn("tasks", result)
        self.assertFalse(result["ai"]["used"])

    def test_analyze_reports_problems(self):
        import webapp
        text = "2026-10-14 09:12 | A | B | 实验三报告下周日交\n坏行\n"
        result = webapp.analyze_text(text, use_ai=False, now=REF)
        self.assertEqual(len(result["problems"]), 1)


class TestExportXlsx(unittest.TestCase):
    """XLSX 导出：真实生成文件、读回验证中文与字段、检查列宽/冻结/筛选。"""

    def _tasks(self):
        messages, _ = read_chat_file(sample_path())
        return build_tasks(messages, now=datetime(2026, 10, 14, 18, 57))

    def test_openpyxl_available(self):
        self.assertTrue(exporter.HAS_OPENPYXL, "本测试需要 openpyxl：pip install openpyxl")

    def test_xlsx_bytes_magic(self):
        data = exporter.to_xlsx(self._tasks())
        self.assertIsInstance(data, bytes)
        self.assertGreater(len(data), 0)
        # xlsx 是 zip 容器，魔数为 PK\x03\x04
        self.assertEqual(data[:2], b"PK")

    def test_xlsx_readable_header_and_chinese(self):
        import openpyxl
        bio = exporter.to_xlsx_stream(self._tasks())
        wb = openpyxl.load_workbook(bio)
        ws = wb.active
        header = [c.value for c in ws[1]]
        self.assertEqual(header, exporter.FIELDS)
        self.assertIn("课程群", header)
        self.assertIn("原始消息", header)
        # 中文课程名 / 中文任务原文 必须正确写入
        chats = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
        self.assertTrue(any(c and "数据结构" in str(c) for c in chats))
        texts = [ws.cell(row=r, column=12).value for r in range(2, ws.max_row + 1)]
        self.assertTrue(any(t and "实验三报告" in str(t) for t in texts))

    def test_xlsx_row_count_and_fields(self):
        import openpyxl
        tasks = self._tasks()
        wb = openpyxl.load_workbook(exporter.to_xlsx_stream(tasks))
        ws = wb.active
        self.assertEqual(ws.max_row, len(tasks) + 1)      # 表头 + 数据
        self.assertEqual(ws.max_column, len(exporter.FIELDS))
        # 截止时间 / 状态 字段与源数据一致
        first = tasks[0]
        row2 = {name: ws.cell(row=2, column=i + 1).value for i, name in enumerate(exporter.FIELDS)}
        self.assertEqual(row2["状态"], first["status"])

    def test_xlsx_freeze_and_filter(self):
        import openpyxl
        wb = openpyxl.load_workbook(exporter.to_xlsx_stream(self._tasks()))
        ws = wb.active
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertIsNotNone(ws.auto_filter.ref)
        self.assertTrue(ws.auto_filter.ref.startswith("A1:"))

    def test_xlsx_column_widths_set(self):
        import openpyxl
        wb = openpyxl.load_workbook(exporter.to_xlsx_stream(self._tasks()))
        ws = wb.active
        for col in ("A", "B", "L"):
            self.assertIsNotNone(ws.column_dimensions[col].width)
            self.assertGreater(ws.column_dimensions[col].width, 5)


class TestCsvBom(unittest.TestCase):
    """CSV 必须带 UTF-8 BOM，且中文可正常读回。"""

    def _tasks(self):
        messages, _ = read_chat_file(sample_path())
        return build_tasks(messages, now=datetime(2026, 10, 14, 18, 57))

    def test_csv_text_has_bom(self):
        text = exporter.to_csv(self._tasks())
        self.assertTrue(text.startswith("\ufeff"))

    def test_csv_bytes_has_bom(self):
        data = exporter.to_csv_bytes(self._tasks())
        self.assertEqual(data[:3], b"\xef\xbb\xbf")

    def test_csv_roundtrip_chinese(self):
        import csv as _csv
        data = exporter.to_csv_bytes(self._tasks())
        text = data.decode("utf-8-sig")
        rows = list(_csv.DictReader(io.StringIO(text)))
        self.assertTrue(rows)
        self.assertIn("课程群", rows[0])
        self.assertTrue(any("数据结构" in r["课程群"] for r in rows))

class TestSmartPaste(unittest.TestCase):
    """智能粘贴：只贴老师消息，自动补时间/发送人/群名。"""

    NOW = datetime(2026, 10, 14, 20, 30)

    def test_bare_message_defaults(self):
        from sentinel.paste_import import parse_smart_paste
        msg, notes = parse_smart_paste("实验三报告下周日晚上12点前交", now=self.NOW)
        self.assertEqual(msg["text"], "实验三报告下周日晚上12点前交")
        self.assertEqual(msg["sender"], "老师")
        self.assertEqual(msg["msg_time"], self.NOW)
        self.assertTrue(notes)  # 明确告知用了默认值

    def test_sender_line_recognized(self):
        from sentinel.paste_import import parse_smart_paste
        msg, notes = parse_smart_paste("发送人: 李老师\n下周三交作业", now=self.NOW)
        self.assertEqual(msg["sender"], "李老师")
        self.assertEqual(msg["text"], "下周三交作业")

    def test_time_and_sender_recognized(self):
        from sentinel.paste_import import parse_smart_paste
        raw = "时间: 2026-10-14 09:12\n发送人: 王助教\n本周五课堂小测"
        msg, notes = parse_smart_paste(raw, now=self.NOW)
        self.assertEqual(msg["msg_time"], datetime(2026, 10, 14, 9, 12))
        self.assertEqual(msg["sender"], "王助教")
        self.assertEqual(notes, [])

    def test_custom_defaults_and_chat(self):
        from sentinel.paste_import import parse_smart_paste, to_standard_line
        msg, _ = parse_smart_paste("月底前交课程设计", default_sender="张老师",
                                   default_chat="软件工程-2026秋", now=self.NOW)
        self.assertEqual(msg["chat"], "软件工程-2026秋")
        line = to_standard_line(msg)
        self.assertIn("|", line)
        # 生成的标准行必须能被标准解析器读回
        parsed, reason = parse_line(line)
        self.assertEqual(reason, "")
        self.assertEqual(parsed["chat"], "软件工程-2026秋")

    def test_empty_returns_problem(self):
        from sentinel.paste_import import parse_smart_paste
        msg, notes = parse_smart_paste("   ", now=self.NOW)
        self.assertIsNone(msg)
        self.assertTrue(notes)

    def test_roundtrip_into_build_tasks(self):
        from sentinel.paste_import import parse_smart_paste, to_standard_line
        msg, _ = parse_smart_paste("实验三报告下周日晚上12点前交", default_chat="数据结构",
                                   now=self.NOW)
        tasks = build_tasks([msg], now=self.NOW)
        self.assertEqual(len(tasks), 1)
        self.assertEqual(tasks[0]["due"].month, 10)
        self.assertEqual(tasks[0]["due"].day, 25)


class TestUploadInspect(unittest.TestCase):
    """上传文件识别：文本 / QQ 加密备份 / 二进制。"""

    def test_plain_text_ok(self):
        from sentinel.upload_inspect import inspect_upload
        data = "2026-10-14 09:12 | 助教-李 | 数据结构 | 实验三报告下周日交\n".encode("utf-8")
        r = inspect_upload(data, "chat.txt")
        self.assertTrue(r["ok"])
        self.assertEqual(r["kind"], "text")
        self.assertIn("助教-李", r["text"])

    def test_gbk_text_ok(self):
        from sentinel.upload_inspect import inspect_upload
        data = "2026-10-14 09:12 | 助教-李 | 数据结构 | 实验三报告下周日交\n".encode("gb18030")
        r = inspect_upload(data, "chat.txt")
        self.assertTrue(r["ok"])
        self.assertIn("助教-李", r["text"])

    def test_qq_backup_identified_with_guide(self):
        import io as _io, zipfile
        from sentinel.upload_inspect import inspect_upload
        buf = _io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            z.writestr("db/1B73A02D1601D5ECCC12125093482ED4", b"SQLite header 3\x00\x04\x00 \x00")
        r = inspect_upload(buf.getvalue(), "2731045816的1个聊天记录.bak")
        self.assertFalse(r["ok"])
        self.assertEqual(r["kind"], "qq_backup")
        self.assertTrue(r["guide"])           # 必须给出操作建议
        self.assertIn("加密", r["error"])

    def test_encrypted_db_identified(self):
        from sentinel.upload_inspect import inspect_upload
        data = b"SQLite header 3\x00\x04\x00 \x00" + b"\x00" * 40 + b"MSG0" + b"\x00" * 200
        r = inspect_upload(data, "nt_msg.db")
        self.assertFalse(r["ok"])
        self.assertEqual(r["kind"], "encrypted_db")

    def test_plain_zip_told_to_unzip(self):
        import io as _io, zipfile
        from sentinel.upload_inspect import inspect_upload
        buf = _io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            z.writestr("a.txt", "hello")
        r = inspect_upload(buf.getvalue(), "records.zip")
        self.assertFalse(r["ok"])
        self.assertEqual(r["kind"], "zip")

    def test_binary_rejected(self):
        from sentinel.upload_inspect import inspect_upload
        data = bytes(range(256)) * 8
        r = inspect_upload(data, "weird.bin")
        self.assertFalse(r["ok"])
        self.assertIn(r["kind"], ("binary", "zip"))

    def test_empty_file(self):
        from sentinel.upload_inspect import inspect_upload
        r = inspect_upload(b"", "empty.txt")
        self.assertFalse(r["ok"])
        self.assertEqual(r["kind"], "empty")


class TestSchedule(unittest.TestCase):
    """结构化时间表：按周一到周日归组 + 分区。"""

    NOW = datetime(2026, 10, 17, 10, 0)   # 2026-10-17 是周六

    def _tasks(self):
        from sentinel.read_chat import parse_lines
        from sentinel.tasks import build_tasks
        text = "\n".join([
            "2026-10-14 09:00 | A | 群 | 实验报告本周日晚上12点前交",
            "2026-10-16 09:00 | A | 群 | 第二次作业下周三交",
            "2026-10-10 09:00 | A | 群 | 第一次作业10月15日24点前交",
        ])
        msgs, _ = parse_lines(text)
        tasks = build_tasks(msgs, now=self.NOW)
        from sentinel.ddl import status_of
        for t in tasks:
            t["status"] = "已作废" if t["superseded"] else status_of(t["due"], self.NOW)
        return tasks

    def test_days_are_monday_to_sunday(self):
        from sentinel.schedule import build_schedule
        sc = build_schedule(self._tasks(), now=self.NOW)
        self.assertEqual(len(sc["days"]), 7)
        self.assertEqual(sc["days"][0]["weekday"], "周一")
        self.assertEqual(sc["days"][6]["weekday"], "周日")
        self.assertEqual(sc["week_start"], "2026-10-12")
        self.assertEqual(sc["week_end"], "2026-10-18")

    def test_today_flagged(self):
        from sentinel.schedule import build_schedule
        sc = build_schedule(self._tasks(), now=self.NOW)
        todays = [d for d in sc["days"] if d["is_today"]]
        self.assertEqual(len(todays), 1)
        self.assertEqual(todays[0]["weekday"], "周六")

    def test_sunday_task_lands_on_sunday(self):
        from sentinel.schedule import build_schedule
        sc = build_schedule(self._tasks(), now=self.NOW)
        sunday = sc["days"][6]
        self.assertTrue(any("实验" in (t.get("text") or "") for t in sunday["items"]))

    def test_overdue_this_week_shown_in_day_column(self):
        """本周内已过期的任务，也要出现在对应的星期列里。"""
        from sentinel.schedule import build_schedule
        tasks = self._tasks()
        sc = build_schedule(tasks, now=self.NOW)
        # 10-15(周四) 那条已过期，应落在周四列（而不是被丢进"已过期"分区就不见）
        thu = sc["days"][3]
        self.assertEqual(thu["weekday"], "周四")
        self.assertTrue(any(t["status"] == "已过期" for t in thu["items"]))
        self.assertGreaterEqual(sc["week_total"], 2)

    def test_buckets_present(self):
        from sentinel.schedule import build_schedule
        sc = build_schedule(self._tasks(), now=self.NOW)
        keys = [b["key"] for b in sc["buckets"]]
        for k in ("已过期", "今天", "明天", "本周", "下周", "更远", "时间未知"):
            self.assertIn(k, keys)
        overdue = [b for b in sc["buckets"] if b["key"] == "已过期"][0]
        self.assertGreaterEqual(overdue["count"], 1)

    def test_superseded_excluded(self):
        from sentinel.schedule import build_schedule
        tasks = self._tasks()
        for t in tasks:
            t["superseded"] = True
        sc = build_schedule(tasks, now=self.NOW)
        self.assertEqual(sc["week_total"], 0)
        self.assertEqual(sc["unscheduled"], 0)

    def test_schedule_csv_export(self):
        from sentinel.schedule import build_schedule
        from sentinel.export import schedule_to_csv
        sc = build_schedule(self._tasks(), now=self.NOW)
        data = schedule_to_csv(sc)
        self.assertTrue(data.startswith(b"\xef\xbb\xbf"))   # BOM
        text = data.decode("utf-8-sig")
        self.assertIn("星期", text.splitlines()[0])
        self.assertIn("实验", text)


class TestQQTranscript(unittest.TestCase):
    """整段 QQ 聊天记录：一次粘贴，自动拆成多条。"""

    NOW = datetime(2026, 10, 17, 10, 0)

    QQ_COPY = """(ᗜ ˰ ᗜ)亓兀元丌: 09-28 08:12:04
        各位同学，提交时间：9月30号周三16:30线代课后。

长缘已尽，故人依旧: 10-13 17:00:47
线代作业 P49 练习2.2 3；4；6；7；10 明天上课的第二节课后上交
"""

    QQ_EXPORT = """消息记录（此消息记录为文本格式，不支持重新导入）
================================================================
消息分组:我加入的群聊
================================================================

2026-10-14 09:12:19 (助教-李)
实验三报告下周日晚上12点前交

2026-10-16 19:01:55 (助教-李)
作业练习2.4：1，4
"""

    def test_copy_format_split_into_messages(self):
        from sentinel.paste_import import parse_smart_paste_batch, to_standard_line
        lines, notes = parse_smart_paste_batch(self.QQ_COPY, default_chat="线代群", now=self.NOW)
        self.assertEqual(len(lines), 2)
        self.assertIn("(ᗜ ˰ ᗜ)亓兀元丌", lines[0])
        self.assertIn("2026-09-28 08:12", lines[0])
        self.assertIn("长缘已尽，故人依旧", lines[1])
        self.assertIn("2026-10-13 17:00", lines[1])
        self.assertIn("2", notes[0])   # 提示里说明识别到 2 条

    def test_export_format_skips_junk_lines(self):
        from sentinel.paste_import import parse_smart_paste_batch
        lines, notes = parse_smart_paste_batch(self.QQ_EXPORT, default_chat="数据结构", now=self.NOW)
        self.assertEqual(len(lines), 2)
        self.assertNotIn("消息分组", "\n".join(lines))
        self.assertNotIn("====", "\n".join(lines))
        self.assertIn("助教-李", lines[0])
        self.assertIn("2026-10-14 09:12", lines[0])

    def test_year_inferred_for_past_month_day(self):
        from sentinel.paste_import import parse_qq_transcript
        # 现在是 2026-10，记录写 09-28 -> 2026；写 12-01（还没到）-> 2025
        msgs = parse_qq_transcript("A: 09-28 08:12:04\n正文1\n\nB: 12-01 08:00:00\n正文2",
                                   chat="群", now=self.NOW)
        self.assertEqual(msgs[0]["msg_time"].year, 2026)
        self.assertEqual(msgs[1]["msg_time"].year, 2025)

    def test_single_message_still_works(self):
        from sentinel.paste_import import parse_smart_paste_batch
        lines, notes = parse_smart_paste_batch("实验三报告下周日晚上12点前交", now=self.NOW)
        self.assertEqual(len(lines), 1)
        self.assertIn("实验三报告下周日晚上12点前交", lines[0])

    def test_standard_format_returned_as_is(self):
        from sentinel.paste_import import parse_smart_paste_batch
        std = "2026-10-14 09:12 | 助教-李 | 数据结构 | 实验三报告下周日交"
        lines, notes = parse_smart_paste_batch(std, now=self.NOW)
        self.assertEqual(lines, [std])

    def test_multiline_body_folded_to_one_line(self):
        from sentinel.paste_import import to_standard_line
        line = to_standard_line({"msg_time": self.NOW, "sender": "A", "chat": "B",
                                 "text": "第一行\n第二行\n第三行"})
        self.assertNotIn("\n", line)
        self.assertIn("第一行 第二行 第三行", line)

    def test_time_word_not_treated_as_sender(self):
        from sentinel.paste_import import parse_qq_transcript
        # "提交时间: 9月30号" 这种不能被当成消息头
        msgs = parse_qq_transcript("A: 09-28 08:12:04\n提交时间: 9月30号周三16:30", chat="群", now=self.NOW)
        self.assertEqual(len(msgs), 1)
        self.assertIn("提交时间", msgs[0]["text"])


class TestSmartPasteApi(unittest.TestCase):
    """智能粘贴的分析流程（不启服务）。"""

    def test_analyze_smart_line(self):
        import webapp
        from sentinel.paste_import import parse_smart_paste, to_standard_line
        msg, _ = parse_smart_paste("实验三报告下周日晚上12点前交", default_chat="数据结构",
                                   now=datetime(2026, 10, 14, 20, 30))
        result = webapp.analyze_text(to_standard_line(msg), use_ai=False,
                                     now=datetime(2026, 10, 14, 20, 30))
        self.assertTrue(result["ok"])
        self.assertEqual(result["summary"]["total"], 1)

class TestAgentCore(unittest.TestCase):
    """Agent 控制层：Agent 编排、Trace、AI 决策。"""

    NOW = datetime(2026, 10, 14, 18, 57)

    def _text(self):
        with open(sample_path(), encoding="utf-8") as f:
            return f.read()

    def test_offline_no_ai_but_full_result(self):
        from sentinel.agent import run_agent
        r = run_agent(self._text(), mode="offline", now=self.NOW)
        self.assertTrue(r["ok"])
        self.assertFalse(r["ai"]["used"])
        self.assertEqual(r["ai"]["mode"], "offline")
        self.assertGreater(r["summary"]["total"], 0)
        # 离线模式必须仍产出完整结果
        self.assertTrue(r["tasks"])

    def test_trace_present_and_ordered(self):
        from sentinel.agent import run_agent
        r = run_agent(self._text(), mode="offline", now=self.NOW)
        trace = r["trace"]
        self.assertGreaterEqual(len(trace), 5)
        self.assertEqual([s["step"] for s in trace], list(range(1, len(trace) + 1)))
        tools = [s["tool"] for s in trace]
        for expected in ("read_chat", "extract", "ddl", "uncertainty"):
            self.assertIn(expected, tools)

    def test_task_has_judge_and_confidence(self):
        from sentinel.agent import run_agent
        r = run_agent(self._text(), mode="offline", now=self.NOW)
        for t in r["tasks"]:
            self.assertIn(t["judge"], ("规则确认", "AI 辅助"))
            self.assertIn(t["confidence_label"], ("高", "中", "低", "未知"))
            self.assertIsNotNone(t["source_index"])

    def test_demo_mode_uses_demo_env(self):
        from sentinel.agent import run_agent
        import os
        os.environ.update({"DEMO_BASE_URL": "http://x/v1", "DEMO_API_KEY": "k", "DEMO_MODEL": "m"})
        try:
            def fake(payload, cfg, timeout):
                return {"choices": [{"message": {"content": "[]"}}]}
            r = run_agent(self._text(), mode="demo", now=self.NOW, transport=fake)
            self.assertEqual(r["ai"]["mode"], "demo")
            self.assertTrue(r["ai"]["enabled"])
        finally:
            for k in ("DEMO_BASE_URL", "DEMO_API_KEY", "DEMO_MODEL"):
                os.environ.pop(k, None)

    def test_ai_failure_falls_back_to_rules(self):
        from sentinel.agent import run_agent
        import os
        os.environ.update({"DEMO_BASE_URL": "http://x/v1", "DEMO_API_KEY": "k", "DEMO_MODEL": "m"})
        try:
            def boom(payload, cfg, timeout):
                from sentinel import ai_helper
                raise ai_helper.AIError("network", "网络断了")
            r = run_agent(self._text(), mode="demo", now=self.NOW, transport=boom)
            self.assertFalse(r["ai"]["used"])         # 失败时不假装用了 AI
            self.assertTrue(r["ai"]["error"])
            self.assertTrue(r["tasks"])               # 仍返回规则结果
        finally:
            for k in ("DEMO_BASE_URL", "DEMO_API_KEY", "DEMO_MODEL"):
                os.environ.pop(k, None)

    def test_empty_input(self):
        from sentinel.agent import run_agent
        r = run_agent("", mode="offline", now=self.NOW)
        self.assertTrue(r["ok"])
        self.assertEqual(r["tasks"], [])


class TestApiConfig(unittest.TestCase):
    """三种 API 模式的配置管理（含 Key 脱敏、内存不落盘）。"""

    def test_offline_disabled(self):
        from sentinel import api_config
        cfg = api_config.effective_config("offline")
        self.assertFalse(cfg["enabled"])
        self.assertEqual(cfg["mode"], "offline")

    def test_custom_save_and_mask(self):
        from sentinel import api_config
        api_config.clear_custom()
        st = api_config.set_custom("https://openrouter.ai/api/v1", "sk-secret-value-9999", "some-model")
        self.assertTrue(st["custom"]["configured"])
        self.assertNotIn("sk-secret-value-9999", json.dumps(st, ensure_ascii=False))
        masked = st["custom"]["key_masked"]
        self.assertIn("*", masked)
        api_config.clear_custom()

    def test_custom_effective(self):
        from sentinel import api_config
        api_config.set_custom("https://a/v1", "k-1234567890", "m1")
        cfg = api_config.effective_config("custom")
        self.assertTrue(cfg["enabled"])
        self.assertEqual(cfg["model"], "m1")
        api_config.clear_custom()

    def test_status_has_no_raw_key(self):
        from sentinel import api_config
        api_config.set_custom("https://a/v1", "sk-rawkey-abcdefgh", "m")
        blob = json.dumps(api_config.status(), ensure_ascii=False)
        self.assertNotIn("sk-rawkey-abcdefgh", blob)
        api_config.clear_custom()


class TestPerVisitorConfig(unittest.TestCase):
    """回归：网页端访客配置按访客隔离（存浏览器本地、随请求带上）。

    关键安全性质：访客传来的 custom 配置绝不回退到服务端环境变量，
    因此访客 A 不会用到作者或访客 B 的 Key；没带配置就回退规则。
    """

    def test_request_custom_no_env_fallback(self):
        from sentinel import api_config, ai_helper
        import os
        # 假装服务端本机配了作者的 Key
        os.environ.update({"BASE_URL": "https://author/v1",
                           "API_KEY": "sk-AUTHOR-ENV-KEY",
                           "MODEL": "author-model"})
        try:
            # 访客没带 api_key（空配置）-> 不能用作者的 Key，必须判定"配置不完整"
            cfg_none = api_config.effective_config("custom", custom={})
            self.assertFalse(cfg_none["enabled"])
            self.assertNotEqual(cfg_none.get("api_key"), "sk-AUTHOR-ENV-KEY")
        finally:
            for k in ("BASE_URL", "API_KEY", "MODEL"):
                os.environ.pop(k, None)

    def test_request_custom_uses_visitor_key(self):
        from sentinel import api_config
        cfg = api_config.effective_config("custom", custom={
            "base_url": "https://visitor/v1", "api_key": "sk-VISITOR", "model": "vm"})
        self.assertTrue(cfg["enabled"])
        self.assertEqual(cfg["api_key"], "sk-VISITOR")
        self.assertEqual(cfg["base_url"], "https://visitor/v1")

    def test_webapp_custom_extraction(self):
        import webapp
        self.assertIsNone(webapp._custom_from_body({"mode": "custom"}))
        got = webapp._custom_from_body({"base_url": "https://b/v1", "api_key": "sk-X", "model": "m"})
        self.assertEqual(got["api_key"], "sk-X")


class TestAgentApi(unittest.TestCase):
    def test_agent_route_via_analyze_helper(self):
        import webapp
        text = "2026-10-14 09:12 | 助教-李 | 数据结构 | 实验三报告下周日晚上12点前交"
        r = webapp.agent_core.run_agent(text, mode="offline", now=datetime(2026, 10, 14, 18, 57))
        self.assertTrue(r["ok"])
        self.assertEqual(r["summary"]["total"], 1)

class TestAIRespectsRules(unittest.TestCase):
    """回归：AI 绝不能覆盖规则已经解析出的时间。"""

    NOW = datetime(2026, 10, 14, 18, 57)

    def test_select_only_picks_unknown_due(self):
        from sentinel import ai_helper
        messages, _ = read_chat_file(sample_path())
        tasks = build_tasks(messages, now=self.NOW)
        picked = ai_helper.select_tasks_for_ai(tasks)
        # 被挑中的必须全部是"规则完全没给出时间"的
        for t in picked:
            self.assertIsNone(t["due"])

    def test_ai_cannot_overwrite_rule_date(self):
        from sentinel import agent as ag
        messages, _ = read_chat_file(sample_path())
        tasks = build_tasks(messages, now=self.NOW)
        # 找一条规则已定时间的任务（例如本周五小测）
        target = [t for t in tasks if t["due"] is not None and "本周五" in t["text"]][0]
        before = target["due"]
        # 伪造一条 AI 建议，试图覆盖它
        ag.apply_ai_suggestions(tasks, [{"id": target["id"], "due": datetime(2026, 1, 1), "reason": "乱给"}], self.NOW)
        self.assertEqual(target["due"], before)          # 未被覆盖
        self.assertEqual(target["origin"], "rule")       # 来源仍是规则

    def test_ai_fills_only_genuinely_unknown(self):
        from sentinel import agent as ag
        messages, _ = read_chat_file(sample_path())
        tasks = build_tasks(messages, now=self.NOW)
        unknown = [t for t in tasks if t["due"] is None and not t["superseded"]]
        if not unknown:
            self.skipTest("样例里没有未知时间的任务")
        tgt = unknown[0]
        ag.apply_ai_suggestions(tasks, [{"id": tgt["id"], "due": datetime(2026, 11, 9, 23, 59), "reason": "x"}], self.NOW)
        self.assertEqual(tgt["due"], datetime(2026, 11, 9, 23, 59))
        self.assertEqual(tgt["origin"], "ai")
        self.assertTrue(tgt["needs_confirm"])

if __name__ == "__main__":
    unittest.main()



